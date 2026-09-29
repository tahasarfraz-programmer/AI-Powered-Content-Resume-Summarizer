"""Turn uploads and URLs into clean plain text."""
import io
import ipaddress
import re
import socket
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX, TXT, or MD file.")

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = unwrap_lines("\n\n".join((page.extract_text() or "") for page in reader.pages))
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="replace")

    text = clean_text(text)
    if len(text) < 30:
        raise ValueError(
            "No readable text found. If this is a scanned PDF, run OCR on it first."
        )
    return text


def unwrap_lines(text: str) -> str:
    """Rejoin PDF hard-wrapped lines: a line that stops mid-sentence and is continued in lowercase."""
    out: list[str] = []
    for ln in text.split("\n"):
        prev = out[-1].rstrip() if out else ""
        if prev and ln.strip() and ln.lstrip()[0].islower() and not re.search(r"[.!?:;][\"')\]]?$", prev):
            out[-1] = prev + " " + ln.lstrip()
        else:
            out.append(ln)
    return "\n".join(out)


def clean_text(text: str) -> str:
    lines = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    out, blank = [], 0
    for ln in lines:
        if ln.strip():
            blank = 0
            out.append(ln)
        else:
            blank += 1
            if blank == 1:
                out.append("")
    return "\n".join(out).strip()


def _is_public_host(host: str) -> bool:
    """Block localhost, private ranges and cloud metadata addresses (SSRF guard)."""
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if (
            ip.is_private
            or ip.is_loopback
            or ip.is_link_local
            or ip.is_reserved
            or ip.is_multicast
        ):
            return False
    return True


async def fetch_url(url: str) -> str:
    current = url.strip()
    if not current.startswith(("http://", "https://")):
        current = "https://" + current

    async with httpx.AsyncClient(timeout=15, follow_redirects=False) as client:
        for _ in range(4):  # follow up to 3 redirects, re-checking each hop
            host = urlparse(current).hostname or ""
            if not _is_public_host(host):
                raise ValueError("That address can't be fetched. Use a public web page.")
            resp = await client.get(current, headers={"User-Agent": "AISummarizerBot/1.0"})
            if resp.is_redirect:
                current = urljoin(current, resp.headers.get("location", ""))
                continue
            resp.raise_for_status()
            break
        else:
            raise ValueError("Too many redirects.")

    ctype = resp.headers.get("content-type", "")
    if "pdf" in ctype:
        return extract_text("page.pdf", resp.content)
    if "html" not in ctype and "text" not in ctype:
        raise ValueError("That link doesn't point to a web page or text file.")

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript"]):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    text = clean_text(root.get_text("\n"))
    if len(text) < 200:
        raise ValueError("Couldn't find enough article text on that page.")
    return text
